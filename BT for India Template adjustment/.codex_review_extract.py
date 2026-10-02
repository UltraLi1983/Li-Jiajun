from pathlib import Path
from docx import Document
import json
import zipfile
from lxml import etree

root = Path('.')
out = {}
for path in sorted(root.glob('*.docx')):
    doc = Document(path)
    item = {"paragraphs": [], "tables": []}
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            item["paragraphs"].append(text)
    for ti, table in enumerate(doc.tables, 1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.replace('\n', ' | ').strip() for cell in row.cells])
        item["tables"].append({"table": ti, "rows": rows})
    out[path.name] = item
    with zipfile.ZipFile(path) as zf:
        xml = etree.fromstring(zf.read("word/document.xml"))
        item["all_document_text"] = [t for t in xml.xpath("//*[local-name()='t']/text()") if t.strip()]
print(json.dumps(out, ensure_ascii=False, indent=2))
