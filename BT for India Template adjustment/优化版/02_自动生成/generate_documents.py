from pathlib import Path
from copy import deepcopy
from datetime import datetime, date
from openpyxl import load_workbook
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
PACK = ROOT / "优化版"
XLSX = PACK / "01_结构化录入" / "印度商务签证_信息采集主表.xlsx"
OUT = PACK / "04_生成结果"
OUT.mkdir(parents=True, exist_ok=True)

wb = load_workbook(XLSX, data_only=False)

def val(v, field):
    if v is None or str(v).strip() == "": return f"<<{field}>>"
    if isinstance(v, (date, datetime)): return v.strftime("%Y-%m-%d")
    return str(v).strip()

def kv(sheet):
    ws=wb[sheet]; result={}
    for r in range(5, ws.max_row+1):
        key=ws.cell(r,1).value
        if key: result[str(key)] = val(ws.cell(r,3).value, str(key))
    return result

a=kv("01_申请人"); c=kv("02_中方公司"); i=kv("03_印度公司"); p=kv("07_Proforma补充")

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)

def base_doc(title=None):
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    styles=d.styles; normal=styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.08
    for sname,size,color in [("Heading 1",15,"17365D"),("Heading 2",12,"244062")]:
        s=styles[sname]; s.font.name="Arial"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    if title:
        q=d.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run(title); r.bold=True; r.font.name="Arial"; r.font.size=Pt(16); r.font.color.rgb=RGBColor(23,54,93)
    return d

def add_kv_table(doc, rows, widths=(2.15,4.75)):
    t=doc.add_table(rows=0, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for label,value in rows:
        cells=t.add_row().cells; cells[0].width=Inches(widths[0]); cells[1].width=Inches(widths[1]); cells[0].text=label; cells[1].text=value
        set_cell_shading(cells[0],"D9EAF7"); cells[0].paragraphs[0].runs[0].bold=True
        for cell in cells: cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t

def save(doc,name):
    path=OUT/name; doc.save(path); print(path)

# 01 Dispatch letter
d=base_doc(); h=d.add_paragraph(); h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; h.add_run(f"Date: {c['dispatch_date']}").bold=True
d.add_paragraph("To: The Embassy of India, Beijing")
d.add_paragraph("Subject: Dispatch Letter in Support of Business Visa Application").runs[0].bold=True
d.add_paragraph("Dear Sir/Madam,")
d.add_paragraph(f"This is to certify that {a['applicant_name']}, holder of passport No. {a['passport_no']} and Chinese ID No. {a['national_id']}, has been employed by {c['cn_company_name']} since {a['employment_start']} and currently serves as {a['job_title']} in the {a['department']} department. The applicant's current annual salary is CNY {a['annual_salary_cny']}.")
d.add_paragraph(f"The applicant will visit {i['in_company_name']}, located at {i['in_company_address']}, India, from {a['trip_start']} to {a['trip_end']} for the purpose of {a['business_purpose']}.")
d.add_paragraph(f"All expenses related to this business trip will be borne by {a['cost_bearer']}. The applicant's position will be retained during the trip, and the company confirms that the applicant will comply with Indian laws and return to China before the authorized stay expires.")
d.add_paragraph("We respectfully request that the appropriate Business Visa be granted.")
d.add_paragraph("Sincerely,")
add_kv_table(d,[("Authorized signatory",c['signatory_name']),("Title / Department",f"{c['signatory_title']} / {c['signatory_department']}"),("Company",c['cn_company_name']),("Telephone",c['signatory_phone']),("Email",c['signatory_email']),("Signature and company seal"," ")])
save(d,"01_Dispatch_Letter.docx")

# 02 Invitation letter
d=base_doc("Invitation Letter")
add_kv_table(d,[("Indian company",i['in_company_name']),("Address",i['in_company_address']),("Inviter",i['inviter_name']),("Title / Department",f"{i['inviter_title']} / {i['inviter_department']}"),("Telephone",i['inviter_mobile']),("Email",i['inviter_email']),("Date",i['invitation_date'])])
d.add_paragraph("To: The Embassy of India, Beijing")
d.add_paragraph(f"We hereby invite {a['applicant_name']}, passport No. {a['passport_no']}, {a['job_title']} of {c['cn_company_name']}, to visit {i['in_company_name']} in India from {a['trip_start']} to {a['trip_end']} for {a['business_purpose']}.")
add_kv_table(d,[("Name",a['applicant_name']),("Passport number",a['passport_no']),("Date of birth",a['dob']),("Position",a['job_title']),("Employer",c['cn_company_name'])])
d.add_paragraph(f"During the visit, all travel, accommodation, insurance and related expenses will be borne by {a['cost_bearer']}. We will ensure that the visitor understands and observes the conditions of stay, complies with Indian laws and leaves India before the authorized stay expires.")
d.add_paragraph("We respectfully request that the appropriate Business Visa be granted. Please contact us if any further information is required.")
d.add_paragraph("Sincerely,")
add_kv_table(d,[("Inviter",i['inviter_name']),("Title",i['inviter_title']),("Company",i['in_company_name']),("Signature / company seal"," ")])
save(d,"02_Invitation_Letter.docx")

# 03 Itinerary
d=base_doc("Business Itinerary")
d.add_paragraph(f"Applicant: {a['applicant_name']}    Passport No.: {a['passport_no']}    Visit: {a['trip_start']} to {a['trip_end']}")
t=d.add_table(rows=1,cols=6); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
headers=["Start date","End date","City","Company / place","Business activities","Flight / hotel / transport"]
for x,hdr in zip(t.rows[0].cells,headers): x.text=hdr; set_cell_shading(x,"17365D"); x.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255); x.paragraphs[0].runs[0].bold=True
ws=wb["04_行程"]; added=0
for r in range(5,25):
    raw=[ws.cell(r,k).value for k in range(1,7)]
    if any(v not in (None,"") for v in raw):
        cells=t.add_row().cells
        for cell,vv,hdr in zip(cells,raw,headers): cell.text=val(vv,hdr)
        added+=1
if not added:
    cells=t.add_row().cells
    for cell,hdr in zip(cells,headers): cell.text=f"<<{hdr}>>"
d.add_paragraph("Applicant's signature: ______________________________        Date: __________________")
save(d,"03_Business_Itinerary.docx")

# 04 Resume
d=base_doc("Personal Resume")
add_kv_table(d,[("Name",a['applicant_name']),("Date of birth",a['dob']),("Address",a['home_address']),("Phone",a['mobile']),("Email",a['email'])])
d.add_heading("Education",level=1)
t=d.add_table(rows=1,cols=4); hdr=["Period","Institution","Degree / major","Notes"]
for cell,x in zip(t.rows[0].cells,hdr): cell.text=x; set_cell_shading(cell,"17365D"); cell.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255)
for r in range(5,17):
    vals=[wb["05_教育经历"].cell(r,k).value for k in range(1,6)]
    if any(v not in (None,"") for v in vals):
        row=t.add_row().cells; data=[f"{val(vals[0],'start')} - {val(vals[1],'end')}",val(vals[2],'institution'),val(vals[3],'degree'),val(vals[4],'notes')]
        for cell,x in zip(row,data): cell.text=x
if len(t.rows)==1:
    row=t.add_row().cells
    for cell,x in zip(row,["<<period>>","<<institution>>","<<degree / major>>","<<notes>>"]): cell.text=x
d.add_heading("Professional Experience",level=1)
t=d.add_table(rows=1,cols=4); hdr=["Period","Company","Position","Responsibilities"]
for cell,x in zip(t.rows[0].cells,hdr): cell.text=x; set_cell_shading(cell,"17365D"); cell.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255)
for r in range(5,23):
    vals=[wb["06_工作经历"].cell(r,k).value for k in range(1,6)]
    if any(v not in (None,"") for v in vals):
        row=t.add_row().cells; data=[f"{val(vals[0],'start')} - {val(vals[1],'end')}",val(vals[2],'company'),val(vals[3],'position'),val(vals[4],'responsibilities')]
        for cell,x in zip(row,data): cell.text=x
if len(t.rows)==1:
    row=t.add_row().cells
    for cell,x in zip(row,["<<period>>","<<company>>","<<position>>","<<responsibilities>>"]): cell.text=x
d.add_heading("Relevant Skills and Experience",level=1); d.add_paragraph(p['skills_experience'])
save(d,"04_Resume.docx")

# 05 Info note
d=base_doc("Visa Information Note")
add_kv_table(d,[("Name",a['applicant_name']),("Indian company",i['in_company_name']),("Chinese company",c['cn_company_name']),("Telephone",a['mobile']),("Email",a['email'])])
save(d,"05_Information_Note.docx")

# 06 Proforma filled copy based on agency DOCX
src=ROOT/"simplified-proforma-Beijing-2026.docx"; d=Document(src)
maps=[
 {0:p['skills_experience'],1:p['years_experience'],2:p['worked_abroad'],4:p['foreign_company'],5:p['foreign_company_contact'],6:p['foreign_project']},
 {0:c['cn_company_name'],1:c['establishment_year'],2:f"{c['cn_company_address']}; {c['cn_company_phone']}; {c['cn_company_email']}",3:c['ownership_type']},
 {2:p['shareholders_5pct'],4:c['turnover_y1'],5:c['turnover_y2'],6:c['turnover_y3'],7:c['listed_status'],8:c['stock_exchange'],9:c['business_sector'],10:c['company_brief'],11:c['india_presence'],12:p['india_entity_name_address'],13:p['india_investment_nature']},
 {0:i['in_company_name'],1:i['in_establishment_year'],2:f"{i['in_company_address']}; {i['in_company_phone']}; {i['in_company_email']}",3:f"{i['in_ownership_type']}; Foreign ownership: {i['foreign_ownership']}; {i['beneficial_owner']}"}
]
for table,m in zip(d.tables,maps):
    for row_idx,text in m.items(): table.rows[row_idx].cells[-1].text=text
# remove obvious template residue
for para in d.paragraphs:
    if para.text.strip() in {"```markdown```","O"}: para.clear()
signatory_values={
    "Name of Signatory(Mandatory):":c['signatory_name'],
    "Designation of the Signatory:":c['signatory_title'],
    "Mobile No.:":c['signatory_mobile'],
    "Telephone Number (Landline):":c['signatory_phone'],
    "Email Address:":c['signatory_email'],
}
for para in d.paragraphs:
    label=para.text.strip()
    if label in signatory_values:
        para.add_run("  "+signatory_values[label]).bold=True
save(d,"06_Proforma_Filled_Copy.docx")
