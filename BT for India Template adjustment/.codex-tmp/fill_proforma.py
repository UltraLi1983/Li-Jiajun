from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path('/Users/lijiajun/Documents/GitHub/python-document/BT for India Template adjustment')
SOURCE = ROOT / 'simplified-proforma-Beijing-2026.docx'
OUTPUT = ROOT / 'outputs' / 'visa_resume_itinerary' / 'BV_Proforma_Li_Jiajun.docx'


def set_cell(cell, text, size=9):
    """Replace only cell content while preserving cell geometry and borders."""
    first = cell.paragraphs[0]
    ppr = deepcopy(first._p.pPr) if first._p.pPr is not None else None
    for child in list(cell._tc):
        if child.tag.endswith('}tcPr'):
            continue
        cell._tc.remove(child)
    p = cell.add_paragraph()
    if ppr is not None:
        if p._p.pPr is not None:
            p._p.remove(p._p.pPr)
        p._p.insert(0, ppr)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def resize_existing_cell_text(cell, size=9):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(size)


doc = Document(SOURCE)

# Details of the applicant
t = doc.tables[0]
set_cell(t.cell(0, 2), 'Certified VDA 6.3 and VDA 6.8 auditor; 20 years of experience in supply chain management, supplier qualification, supplier development, process auditing and supply chain risk mitigation.')
set_cell(t.cell(1, 2), '20 years')
set_cell(t.cell(2, 2), 'No')
set_cell(t.cell(4, 2), 'N/A')
set_cell(t.cell(5, 2), 'N/A')
set_cell(t.cell(6, 2), 'N/A')

# Details of the applicant's current employer
t = doc.tables[1]
set_cell(t.cell(0, 2), 'Audi China Enterprise Management Co., Ltd.')
set_cell(t.cell(1, 2), '2009')
set_cell(t.cell(2, 2), 'Tower 1, No. 12 Qishengzhong Street, Chaoyang District, Beijing 100028, P.R. China\nTelephone and company email: To be confirmed by employer')
set_cell(t.cell(3, 2), 'Private enterprise (wholly foreign-owned)')

t = doc.tables[2]
set_cell(t.cell(0, 2), '')
set_cell(t.cell(1, 2), 'N/A')
set_cell(t.cell(2, 2), 'N/A')
set_cell(t.cell(2, 1), 'If JV, ownership breakdown:', 9)
set_cell(t.cell(3, 2), 'AUDI AG, Germany - 100%')
set_cell(t.cell(4, 2), 'FY2023: To be confirmed by employer')
set_cell(t.cell(5, 2), 'FY2024: To be confirmed by employer')
set_cell(t.cell(6, 2), 'FY2025: To be confirmed by employer')
set_cell(t.cell(7, 2), 'Non-Listed')
set_cell(t.cell(8, 2), 'N/A')
set_cell(t.cell(8, 1), 'If listed, name of stock exchange:', 9)
set_cell(t.cell(9, 2), 'Automotive industry; corporate management, strategy, research and development, purchasing, marketing, sales and technical support services.')
set_cell(t.cell(10, 1), 'Please share a brief note about the Company.\n\nAudi China Enterprise Management Co., Ltd. was established in Beijing in 2009 and is wholly owned by AUDI AG. It coordinates Audi business activities in China and provides corporate management, strategy, research and development, purchasing, marketing, sales and technical support services.')
set_cell(t.cell(11, 1), "Details of the Company's presence in India through subsidiaries, permanent establishments, joint ventures or beneficial interests.\n\nN/A for Audi China Enterprise Management Co., Ltd., subject to employer confirmation.")
set_cell(t.cell(12, 2), 'N/A')
set_cell(t.cell(13, 2), 'N/A')

# Details of the company in India
t = doc.tables[3]
set_cell(t.cell(0, 2), 'ANAND NVH PRODUCTS PVT. LTD.')
set_cell(t.cell(1, 2), '1988 (operations under the current name commenced in 2002)')
set_cell(t.cell(2, 2), 'Plot No. 33, Sector 35, HSIIDC, Gurugram, Haryana 122001, India\nTel.: +91-124-4030580; Email: lalit.koul@anandnvh.com')
set_cell(t.cell(3, 1), 'Nature of ownership: Private company\nForeign ownership: To be confirmed by Indian company')
set_cell(t.cell(3, 2), 'Beneficial owner nationality and ownership percentage: To be confirmed by Indian company')

# Remove stray conversion artifacts above the signature-details block.
for p in doc.paragraphs:
    if p.text.strip() in {'```markdown```', 'O'}:
        p.text = ''

# Per user instruction, keep the entire authorized-signatory and seal block blank.
for p in doc.paragraphs:
    if p.text.strip().startswith(('Name of Signatory', 'Designation of the Signatory', 'Mobile No.', 'Telephone Number (Landline)', 'Email Address', 'Official Seal')):
        # Labels are preserved; no value is appended.
        pass

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
