# -*- coding: utf-8 -*-
"""几何与文字溢出校验"""
import math
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
prs = Presentation('创始团队合作与治理结构建议.pptx')
W, H = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
print(f'page: {W:.2f} x {H:.2f} in\n')

def is_cjk(ch):
    return ord(ch) > 0x2E80

def est_lines(text, size_pt, width_in):
    cw = sum(size_pt/72 * (1.0 if is_cjk(c) else 0.55) for c in text)
    return max(1, math.ceil(cw / width_in))

def est_height(lines_spec, width_in):
    """估算多段文本总高(in): lines_spec = [(text,size_pt)]"""
    total = 0.0
    for text, size in lines_spec:
        n = est_lines(text, size, width_in)
        total += n * size * 1.28 / 72 + 5 / 72   # 行高 + space_after 5pt
    return total

issues = []
for si, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        x, y = sh.left / EMU_IN, sh.top / EMU_IN
        w, h = sh.width / EMU_IN, sh.height / EMU_IN
        if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
            issues.append(f'slide{si}: OUT-OF-PAGE {sh.shape_type} ({x:.2f},{y:.2f},{w:.2f},{h:.2f})')
        # 文本框文字高度估算
        if sh.has_text_frame and sh.text_frame.text:
            wd = w - (sh.text_frame.margin_left + sh.text_frame.margin_right) / EMU_IN
            specs = []
            for p in sh.text_frame.paragraphs:
                t = ''.join(r.text for r in p.runs)
                if t:
                    size = (p.runs[0].font.size.pt if p.runs[0].font.size else 12)
                    specs.append((t, size))
            eh = est_height(specs, wd)
            if eh > h + 0.08:
                issues.append(f'slide{si}: TEXT-OVERFLOW est {eh:.2f}in > box {h:.2f}in :: {specs[0][0][:24]}...')

print('ISSUES:' if issues else 'NO ISSUES')
for i in issues: print(' -', i)
