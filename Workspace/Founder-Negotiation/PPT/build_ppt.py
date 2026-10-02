# -*- coding: utf-8 -*-
"""三页 PPT：创始团队合作与治理结构建议（带图片）"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE   = os.path.dirname(os.path.abspath(__file__))
IMG_KN = '/Users/lijiajun/Documents/GitHub/python-document/审核系统+运营系统架构/docs/architecture/知识中台<->运营系统（+AI模型）.png'
IMG_FULL = '/Users/lijiajun/Documents/GitHub/python-document/审核系统+运营系统架构/docs/architecture/全运营过程系统-功能架构总览-v4.png'

NAVY   = RGBColor(0x16, 0x32, 0x4F)
BLUE   = RGBColor(0x2E, 0x6D, 0xA4)
TEAL   = RGBColor(0x00, 0xA9, 0x9D)
ORANGE = RGBColor(0xF5, 0xA6, 0x23)
GRAY   = RGBColor(0x8A, 0x94, 0xA6)
LIGHT  = RGBColor(0xEE, 0xF3, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x33, 0x3A, 0x45)
FONT   = 'Hiragino Sans GB'

EMU_IN = 914400

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w or 1.0)
    sp.shadow.inherit = False
    return sp

def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=4, wrap=True):
    """lines: list of (text, size, bold, color) 或 (text, size, bold, color, align)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(lines):
        text, size, bold, color = item[:4]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = item[4] if len(item) > 4 else align
        p.space_after = Pt(space_after)
        p.line_spacing = 1.12
        r = p.add_run(); r.text = text
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color
    return tb

def header(slide, idx, title, subtitle):
    add_rect(slide, 0, 0, 13.333, 0.07, fill=TEAL)
    add_text(slide, 0.55, 0.28, 9.6, 0.7,
             [(title, 27, True, NAVY)])
    add_text(slide, 0.57, 0.92, 6.5, 0.32,
             [(subtitle, 12.5, False, BLUE)])
    add_text(slide, 12.0, 0.45, 0.8, 0.4,
             [(f'{idx} / 3', 13, True, GRAY)], align=PP_ALIGN.RIGHT)

def footer(slide):
    add_text(slide, 0.55, 7.14, 12.2, 0.28,
             [('本文件为可供双方讨论的中性版本，不构成有约束力的要约或法律意见', 9, False, GRAY)])

def bullet_card(slide, x, y, w, h, title, items, title_color=NAVY, fill=LIGHT, size=11.5):
    add_rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    lines = [(title, 14.5, True, title_color)]
    for it in items:
        lines.append(('•  ' + it, size, False, DARK))
    add_text(slide, x + 0.25, y + 0.16, w - 0.5, h - 0.3, lines, space_after=5)

# ============ 第 1 页：合作目标与模式选择 ============
s1 = prs.slides.add_slide(BLANK)
header(s1, 1, '创始团队合作与治理结构建议', '01 合作目标与模式选择')

# 左栏：合作目标卡片
bullet_card(s1, 0.55, 1.30, 7.10, 1.18, '合作目标',
            ['共建工业 AI 软件平台：工业 Know-how + 体系标准 + AI 技术；',
             '兼顾资本安全、核心团队长期激励与后续融资能力。'], size=10.5)

# 左栏：模式对比图
ratio = 753 / 1486
img_w, img_h = 7.10, 7.10 * ratio
s1.shapes.add_picture(os.path.join(HERE, 'model_compare.png'), Inches(0.55), Inches(2.68),
                      width=Inches(img_w), height=Inches(img_h))

# 左栏底部结论
add_text(s1, 0.55, 6.40, 7.10, 0.6,
         [('关键提醒：不建议将“联合创始人的责任”与“职业经理人的收益和权限”混搭组合。', 12, True, ORANGE)])

# 右栏：工业 AI 平台架构竖图
kn = Image.open(IMG_KN)
kr = kn.size[1] / kn.size[0]
kh, kw = 5.42, 5.42 / kr          # 等比缩放，控制高度
kx = 8.45
s1.shapes.add_picture(IMG_KN, Inches(kx), Inches(1.32), width=Inches(kw), height=Inches(kh))
add_rect(s1, kx - 0.06, 1.26, kw + 0.12, kh + 0.12, fill=None, line=RGBColor(0xD5, 0xDE, 0xE8), line_w=1.0)
add_text(s1, 8.45, 6.85, 4.4, 0.28, [('工业 AI 软件平台架构示意（知识中台 ↔ 运营系统 + AI）', 9, False, GRAY)], align=PP_ALIGN.CENTER)
footer(s1)

# ============ 第 2 页：经济结构与股权方案 ============
s2 = prs.slides.add_slide(BLANK)
header(s2, 2, '经济结构与股权方案', '02 股权、归属与融资稀释')

# 左栏：股权对比图
e = Image.open(os.path.join(HERE, 'equity_chart.png'))
er = e.size[1] / e.size[0]
ew, eh = 7.10, 7.10 * er
s2.shapes.add_picture(os.path.join(HERE, 'equity_chart.png'), Inches(0.55), Inches(1.32),
                      width=Inches(ew), height=Inches(eh))
card_y2 = 1.32 + eh + 0.20
bullet_card(s2, 0.55, card_y2, 7.10, 7.12 - card_y2, '方案要点',
            ['ESOP 按 18–24 个月招聘计划确定，初步可讨论 10% 并按需扩充，避免无岗位预算一次预留 25%；',
             '联合创始人模式建议谈判区间 55/35/10 或 51/39/10；',
             '若坚持 70/5/25，须按职业经理人模式补足薪酬、奖金、动态期权、经营授权与无因解聘保护。'], size=10.5)

# 右栏：Vesting 图 + 要点
v = Image.open(os.path.join(HERE, 'vesting_chart.png'))
vr = v.size[1] / v.size[0]
vw, vh = 4.80, 4.80 * vr
s2.shapes.add_picture(os.path.join(HERE, 'vesting_chart.png'), Inches(8.05), Inches(1.32),
                      width=Inches(vw), height=Inches(vh))
card_y2r = 1.32 + vh + 0.20
bullet_card(s2, 8.05, card_y2r, 4.80, 7.12 - card_y2r, '归属与退出',
            ['四年 Vesting + 一年 Cliff，此后按月归属；',
             '成立前已完成的验证工作可讨论起算信用；',
             '未实缴资本对应股权不应无条件全部生效；',
             '明确离职、无因解聘、违约、死亡失能与控制权变更处理；',
             '约定已/未归属股份的回购权、价格、程序与支付。'], size=10.5)
footer(s2)

# ============ 第 3 页：治理结构与推进路径 ============
s3 = prs.slides.add_slide(BLANK)
header(s3, 3, '治理结构与推进路径', '03 权责对等、风险约束与落地节奏')

# 左栏：全运营架构图
f = Image.open(IMG_FULL)
fr = f.size[1] / f.size[0]
fw, fh = 7.35, 7.35 * fr
s3.shapes.add_picture(IMG_FULL, Inches(0.55), Inches(1.32), width=Inches(fw), height=Inches(fh))
add_rect(s3, 0.49, 1.26, fw + 0.12, fh + 0.12, fill=None, line=RGBColor(0xD5, 0xDE, 0xE8), line_w=1.0)
add_text(s3, 0.55, 6.35, 7.35, 0.35, [('全运营过程系统功能架构总览——公司业务全景与产品基础', 10, False, GRAY)], align=PP_ALIGN.CENTER)

# 左栏底部：治理关键词
add_text(s3, 0.55, 6.78, 7.35, 0.35,
         [('控制权、经济权与岗位权责分别表达；以治理规则覆盖融资、稀释、ESOP 扩池与董事会变化。', 11, True, ORANGE)])

# 右栏：治理要点
bullet_card(s3, 8.05, 1.32, 4.80, 2.52, '治理结构',
            ['CEO：董事会批准的战略与预算内，享有日常经营决策权；',
             '董事会：审议年度预算、融资、重大合同、核心高管与资本事项；',
             '保留事项：章程修改、增资减资、并购清算、核心 IP 处置、重大关联交易；',
             '月度经营报告、季度董事会、财务信息权、关联交易审批与僵局解决机制。'], size=10.5)

# 右栏：7 步推进
steps = ['① 确认合作模式（联合创始人 / 职业经理人）',
         '② 交换并核实出资、资源、知识产权、全职投入证明',
         '③ 确认 24 个月预算、里程碑与招聘计划',
         '④ 完成股权与融资稀释模型',
         '⑤ 签署原则性 Term Sheet',
         '⑥ 分别征询法律与税务意见',
         '⑦ 签署股东协议、章程、投资协议、激励与知识产权文件']
card_y = 3.98
add_rect(s3, 8.05, card_y, 4.80, 3.14, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
lines = [('推进路径（7 步）', 14.5, True, NAVY)] + [(s, 10.5, False, DARK) for s in steps]
add_text(s3, 8.30, card_y + 0.14, 4.30, 2.9, lines, space_after=3.5)
footer(s3)

out = os.path.join(HERE, '创始团队合作与治理结构建议.pptx')
prs.save(out)
print('saved ->', out)
