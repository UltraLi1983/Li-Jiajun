#!/usr/bin/env python3
"""Set VDA audit tool to 10-point scoring per question."""

from __future__ import annotations

import re
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

SCORE_LIST = '"n.e.,0,4,6,8,10"'
QUESTION_RE = re.compile(r"^\d+\.\d+\.\d+$")
FORMULA_8_RE = re.compile(r"\(D(\d+)\*8\)")


def iter_score_rows(ws):
    """Score is in F(row+1) when question id is in A(row)."""
    for r in range(1, ws.max_row):
        a = ws[f"A{r}"].value
        if not (a and isinstance(a, str) and QUESTION_RE.match(a.strip())):
            continue
        score_row = r + 1
        cell = ws[f"F{score_row}"]
        if cell.data_type == "f":
            continue
        yield score_row, cell


def main() -> None:
    src = Path(__file__).resolve().parent / "VDA_6.8_Complete_Audit_Tool.xlsm"
    backup = src.with_suffix(".xlsm.bak")
    if not backup.exists():
        copy2(src, backup)

    wb = load_workbook(src, keep_vba=True)
    modules = [n for n in wb.sheetnames if re.match(r"^L\d", n)]

    changed_scores = 0
    changed_formulas = 0

    for name in modules:
        ws = wb[name]
        score_refs: list[str] = []

        for score_row, cell in iter_score_rows(ws):
            score_refs.append(f"F{score_row}")
            if cell.value in ("n.e.", None, ""):
                continue
            if isinstance(cell.value, (int, float)) or str(cell.value).isdigit():
                cell.value = 10
                changed_scores += 1

        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and cell.value:
                    v = str(cell.value)
                    new_v = FORMULA_8_RE.sub(r"(D\1*10)", v)
                    if new_v != v:
                        cell.value = new_v
                        changed_formulas += 1

        if ws.data_validations:
            ws.data_validations.dataValidation = [
                dv
                for dv in ws.data_validations.dataValidation
                if not (dv.type == "list" and dv.sqref and "F" in str(dv.sqref))
            ]

        if score_refs:
            dv = DataValidation(type="list", formula1=SCORE_LIST, allow_blank=True)
            for ref in score_refs:
                dv.add(ref)
            ws.add_data_validation(dv)

    wb["Datasheet"]["B5"] = (
        "Scoring: each question max 10 pts (VDA: 0 / 4 / 6 / 8 / 10); "
        "n.e. = not evaluated"
    )
    wb["Evaluationsheet"]["K5"] = "Per-question max: 10 points"

    wb.save(src)
    print(f"Saved: {src}")
    print(f"Backup: {backup}")
    print(f"Modules: {len(modules)}")
    print(f"Scores set to 10: {changed_scores}")
    print(f"Formulas D*n*8 -> D*n*10: {changed_formulas}")


if __name__ == "__main__":
    main()
