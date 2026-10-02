#!/usr/bin/env python3
"""Fill Excel D/E columns from VDA 6.8 PDF pages 52-158."""

from __future__ import annotations

import re
from pathlib import Path
from shutil import copy2

import fitz
from openpyxl import load_workbook

PDF_PATH = Path(__file__).parent / "VDA_Band_06.8_1. Auflage 2024_Englisch.pdf"
XLSM_PATH = Path(__file__).parent / "VDA_6.8_Complete_Audit_Tool.xlsm"
PAGE_START = 52
PAGE_END = 158  # L7.6.2 full text starts on printed page 158


def sanitize_excel(text: str) -> str:
    if not text:
        return text
    text = text.replace("\u00ad", "")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return text


def load_pdf_text() -> str:
    doc = fitz.open(PDF_PATH)
    text = ""
    for i in range(PAGE_START - 1, PAGE_END):
        text += doc[i].get_text() + "\n"
    text = re.sub(
        r"Dokument wurde bereitgestellt.*?bestimmt\.",
        "",
        text,
        flags=re.DOTALL,
    )
    text = re.sub(
        r"Nur zur internen Verwendung.*?\.",
        "",
        text,
        flags=re.DOTALL,
    )
    text = re.sub(r"\n\d+\nQuestionnaire\n", "\n", text)
    return text


def trim_block(block: str, l_num: str) -> str:
    next_l = str(int(l_num) + 1)
    stops = [
        r"\nProcess element L\d",
        r"\nFigure \d",
        r"\nQuestionnaire\b",
        r"\n>>\s*\n>>",
        rf"\nL{next_l}\.",
        rf"\n{next_l}\.\d",
    ]
    if l_num == "1":
        stops.extend([r"\nPlanning of logistical processes", r"\nL2\.1\.1"])
    cut = len(block)
    for pat in stops:
        m = re.search(pat, block, re.IGNORECASE)
        if m:
            cut = min(cut, m.start())
    return block[:cut].strip()


def normalize_lines(body: str) -> list[str]:
    body = re.sub(r"(\w)-\n(\w)", r"\1\2", body)
    merged: list[str] = []
    for raw in body.split("\n"):
        line = raw.strip()
        if not line:
            continue
        if (
            line.startswith("•")
            or line.startswith("-")
            or re.match(r"^\d+\.\s", line)
        ):
            merged.append(line)
        elif merged:
            merged[-1] += " " + line
        else:
            merged.append(line)
    return merged


def split_numbered_sections(body: str) -> list[tuple[str, str]]:
    lines = normalize_lines(body)
    sections: list[tuple[str, list[str]]] = []
    current_num = None
    current_lines: list[str] = []

    for line in lines:
        m = re.match(r"^(\d+)\.\s*(.*)$", line)
        if m:
            if current_num is not None:
                sections.append((current_num, current_lines))
            current_num = m.group(1)
            rest = m.group(2).strip()
            current_lines = [rest] if rest else []
        else:
            if current_num is None:
                if not sections:
                    sections.append(("0", []))
                sections[-1][1].append(line)
            else:
                current_lines.append(line)
    if current_num is not None:
        sections.append((current_num, current_lines))

    return [(num, "\n".join(lns).strip()) for num, lns in sections]


def parse_question_body(body: str) -> tuple[str, str]:
    if "Minimum requirements relevant for evaluation" in body:
        body = body.split("Minimum requirements relevant for evaluation", 1)[1]
    if "Examples for implementation" in body:
        body = body.split("Examples for implementation", 1)[1]
    body = body.strip()

    sections = split_numbered_sections(body)
    if not sections:
        return "", body

    if len(sections) == 1 and sections[0][0] == "0":
        content = sections[0][1]
        lines = [ln for ln in content.split("\n") if ln.strip()]
        prose = [ln for ln in lines if not (ln.startswith("•") or ln.startswith("-"))]
        bullets = [ln for ln in lines if ln.startswith("•") or ln.startswith("-")]
        if bullets and prose:
            return "\n".join(prose), "\n".join(bullets)
        if bullets:
            return "", "\n".join(bullets)
        return content, ""

    min_parts: list[str] = []
    ex_parts: list[str] = []

    for idx, (num, content) in enumerate(sections):
        if num == "0":
            if content:
                min_parts.append(content)
            continue
        lines = content.split("\n")
        bullets = [ln for ln in lines if ln.startswith("•") or ln.startswith("-")]
        non_bullets = [ln for ln in lines if not (ln.startswith("•") or ln.startswith("-"))]

        is_last = idx == len(sections) - 1
        if is_last and bullets:
            if non_bullets:
                min_parts.append(f"{num}. " + "\n".join(non_bullets))
            ex_parts.extend(bullets)
        else:
            block_lines: list[str] = []
            if non_bullets:
                block_lines.append(f"{num}. " + "\n".join(non_bullets))
            elif num != "0":
                block_lines.append(f"{num}.")
            block_lines.extend(bullets)
            min_parts.append("\n".join(block_lines))

    return "\n\n".join(min_parts).strip(), "\n\n".join(ex_parts).strip()


def extract_all_questions(text: str) -> dict[str, dict[str, str]]:
    pattern = re.compile(
        r"L\.?(?P<l>\d)\.(?P<p>\d+)\.(?P<q>\d+)\*?\s*[\n\t]",
        re.MULTILINE,
    )
    records: dict[str, dict[str, str]] = {}
    for m in pattern.finditer(text):
        start = m.end()
        nxt = pattern.search(text, start)
        end = nxt.start() if nxt else len(text)
        block = trim_block(text[start:end], m.group("l"))
        key = f"{m.group('l')}.{m.group('p')}.{m.group('q')}"
        d, e = parse_question_body(block)
        records[key] = {"d": d, "e": e}
    return records


def format_cell(text: str) -> str | None:
    if not text:
        return None
    text = sanitize_excel(text)
    text = text.replace("\t", " ")
    text = re.sub(r" +", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() or None


def main() -> None:
    backup = XLSM_PATH.with_suffix(".xlsm.bak_pdf_fill")
    if not backup.exists():
        copy2(XLSM_PATH, backup)

    pdf_data = extract_all_questions(load_pdf_text())
    wb = load_workbook(XLSM_PATH, keep_vba=True)

    updated = 0
    missing: list[str] = []

    for name in wb.sheetnames:
        if not (name.startswith("L") and len(name) > 2 and name[1].isdigit()):
            continue
        ws = wb[name]
        for r in range(1, ws.max_row):
            a = ws[f"A{r}"].value
            if not (a and isinstance(a, str) and re.match(r"^\d+\.\d+\.\d+$", a.strip())):
                continue
            qid = a.strip()
            if qid not in pdf_data:
                missing.append(f"{name}:{qid}")
                continue
            row = r + 1
            d = format_cell(pdf_data[qid]["d"])
            e = format_cell(pdf_data[qid]["e"])
            ws[f"D{row}"] = d
            ws[f"E{row}"] = e
            updated += 1

    wb.save(XLSM_PATH)
    print(f"Saved: {XLSM_PATH}")
    print(f"Backup: {backup}")
    print(f"PDF entries: {len(pdf_data)}")
    print(f"Updated rows: {updated}")
    print(f"Missing: {missing}")


if __name__ == "__main__":
    main()
